#!/usr/bin/env python3
"""Таблица сортов <-> файлы данных.
   python3 tools/xlsx_tool.py export [файл.xlsx]   — выгрузить все сорта в таблицу
   python3 tools/xlsx_tool.py import файл.xlsx            — загрузить полную таблицу (если в ней не хватает сортов из базы, импорт остановится)
   python3 tools/xlsx_tool.py import файл.xlsx --merge    — добавить/обновить только строки из файла, остальные сорта остаются
   python3 tools/xlsx_tool.py import файл.xlsx --replace  — база станет ровно такой, как таблица (сорта, которых нет в таблице, удаляются)
   В таблице: пункты «Особенности» и «Выращивание» пишутся в одной ячейке, каждый с новой строки."""
import json, os, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from validate import validate, load_all, ROOT, FILES

CATS = {'apple': 'Яблоня', 'pear': 'Груша'}
RCATS = {v: k for k, v in CATS.items()}
# (заголовок, поле, ширина)
COLS = [('id', 'id', 7), ('Культура', 'cat', 11), ('Название', 'name', 24), ('Год кратко', 'y', 11),
        ('Год появления', 'year', 18), ('Пояснение к году', 'yearNote', 36), ('Происхождение', 'origin', 34),
        ('Родители', 'parents', 28), ('Срок созревания', 'season', 24), ('Зимостойкость', 'hardy', 22),
        ('Лёжкость', 'keep', 26), ('Вкус', 'taste', 24), ('Особенности (каждый пункт с новой строки)', 'facts', 60),
        ('Выращивание (каждый пункт с новой строки)', 'care', 44), ('Интересный факт', 'fun', 44),
        ('Устойчив к парше (да/нет)', 'scab', 14), ('Источники', 'src', 44),
        ('Commons: имя файла', 'commons.file', 24), ('Commons: автор', 'commons.author', 18), ('Commons: лицензия', 'commons.license', 14),
        ('Статус проверки (проверено/черновик)', 'status', 18), ('Проверено (ГГГГ-ММ)', 'checked', 14)]

def cell_value(v, field):
    if field == 'cat': return CATS.get(v.get('cat'), v.get('cat'))
    if field in ('facts', 'care'): return '\n'.join(v.get(field, []))
    if field == 'scab': return 'да' if v.get('scab') == 'res' else 'нет'
    if field == 'status': return 'проверено' if v.get('status') == 'verified' else 'черновик'
    if field.startswith('commons.'): return (v.get('commons') or {}).get(field.split('.')[1], '')
    return v.get(field, '')

def export(path):
    wb = Workbook(); ws = wb.active; ws.title = 'Сорта'
    ws.append([c[0] for c in COLS])
    for _, v in load_all():
        if v['cat'] in CATS: ws.append([cell_value(v, f) for _, f, _ in COLS])
    head = PatternFill('solid', fgColor='2C4A3E')
    for i, (title, _, width) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=i); c.font = Font(bold=True, color='FFFFFF'); c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'D2'; ws.row_dimensions[1].height = 42
    n = ws.max_row + 300
    d1 = DataValidation(type='list', formula1='"Яблоня,Груша"', allow_blank=True); ws.add_data_validation(d1); d1.add(f'B2:B{n}')
    d2 = DataValidation(type='list', formula1='"да,нет"', allow_blank=True); ws.add_data_validation(d2); d2.add(f'P2:P{n}')
    info = wb.create_sheet('Инструкция')
    for line in [
        'Как пользоваться таблицей',
        '1. Каждая строка — один сорт. Новый сорт добавьте строкой внизу: id можно оставить пустым, он присвоится сам.',
        '2. Обязательно заполнить: Культура, Название, Год кратко, Год появления, Происхождение, Срок созревания, Зимостойкость, Лёжкость, Вкус, Источники и минимум 2 пункта в «Особенностях».',
        '3. Если данных нет, пишите «Уточнить». Не придумывайте: каждая цифра должна иметь источник.',
        '4. В «Срок созревания» должно быть слово «лет…», «осен…» или «зим…» (например, «Осенне-зимний»): по нему работает фильтр.',
        '5. Если источники расходятся по году, главный год пишите в «Год появления», а расхождение в «Пояснение к году».',
        '6. Фото с Commons: заполните три колонки Commons (имя файла, автор, лицензия). Другие фото добавляйте только с разрешения автора.',
        '7. Пришлите заполненный файл в чат: данные проверятся, сайт соберётся заново.']:
        info.append([line])
    info.column_dimensions['A'].width = 130; info['A1'].font = Font(bold=True, size=14)
    for r in info.iter_rows():
        for c in r: c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(path); print('Таблица сохранена:', path, '| сортов:', ws.max_row - 1)

def next_id(existing, prefix):
    n = max([int(i[1:]) for i in existing if i.startswith(prefix)] or [0]) + 1
    return f'{prefix}{n}'

def import_xlsx(path, mode='safe'):
    wb = load_workbook(path, data_only=True); ws = wb['Сорта']
    header = [c.value for c in ws[1]]
    idx = {f: header.index(t) for t, f, _ in COLS if t in header}
    old = {v['id']: v for _, v in load_all() if v['cat'] in CATS}
    used = set(old); result = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(x not in (None, '') for x in row): continue
        g = lambda f: ('' if row[idx[f]] is None else str(row[idx[f]]).strip()) if f in idx else ''
        cat = RCATS.get(g('cat'), g('cat'))
        vid = g('id') or next_id(used, 'a' if cat == 'apple' else 'p'); used.add(vid)
        v = dict(old.get(vid, {}))                     # неизвестные поля (фото, галерея) сохраняются
        v.update(id=vid, cat=cat)
        for f in ('name', 'y', 'year', 'yearNote', 'origin', 'parents', 'season', 'hardy', 'keep', 'taste', 'fun', 'src'):
            v[f] = g(f)
        for f in ('facts', 'care'):
            v[f] = [x.strip() for x in g(f).split('\n') if x.strip()]
        v['scab'] = 'res' if g('scab').lower() in ('да', 'yes', 'res', '1') else ''
        st = g('status').lower()
        ck = g('checked') or old.get(vid, {}).get('checked', '')
        if ck: v['checked'] = ck
        else: v.pop('checked', None)
        v['status'] = 'verified' if st in ('проверено', 'verified') else 'draft' if st in ('черновик', 'draft') else old.get(vid, {}).get('status', 'draft')
        cf = g('commons.file')
        if cf: v['commons'] = {'file': cf, 'author': g('commons.author'), 'license': g('commons.license')}
        else: v.pop('commons', None)
        for k in ('yearNote', 'parents', 'fun'):
            if not v.get(k): v.pop(k, None)
        result.append(v)
    sheet_ids = {v['id'] for v in result}
    missing = [i for i in old if i not in sheet_ids]
    if mode == 'merge':
        result = [v for v in old.values() if v['id'] not in sheet_ids] + result
    elif mode == 'safe' and missing:
        print(f'ОСТАНОВЛЕНО: в таблице нет {len(missing)} сортов, которые есть в базе: ' + ', '.join(f'{i} ({old[i].get("name")})' for i in missing[:10]) + (' …' if len(missing) > 10 else ''))
        print('Это похоже на таблицу только с новыми сортами. Используйте --merge, чтобы добавить их к базе, или --replace, если удаление нужно.')
        sys.exit(1)
    errors, warnings = validate([(FILES.get(v['cat']), v) for v in result])
    for w in warnings: print('  предупреждение:', w)
    if errors:
        for e in errors: print('  ОШИБКА:', e)
        print('Импорт остановлен: данные не записаны.'); sys.exit(1)
    for cat, rel in FILES.items():
        if cat not in CATS: continue          # декоративные растения таблица не затрагивает
        json.dump([v for v in result if v['cat'] == cat], open(os.path.join(ROOT, rel), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'Импорт выполнен: {len(result)} сортов записано.')

if __name__ == '__main__':
    if len(sys.argv) >= 2 and sys.argv[1] == 'export': export(sys.argv[2] if len(sys.argv) > 2 else 'sorta.xlsx')
    elif len(sys.argv) >= 3 and sys.argv[1] == 'import':
        flag = sys.argv[3] if len(sys.argv) > 3 else ''
        import_xlsx(sys.argv[2], 'merge' if flag == '--merge' else 'replace' if flag == '--replace' else 'safe')
    else: print(__doc__)
